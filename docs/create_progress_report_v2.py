import os
import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

def create_report():
    doc = docx.Document()

    # Thiết lập lề trang tiêu chuẩn 2.0 cm
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style mặc định
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(33, 33, 33)

    # 1. TIÊU ĐỀ BÁO CÁO
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI GR2")
    run_title.bold = True
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(15)
    run_title.font.color.rgb = RGBColor(18, 52, 86)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Đề tài: Xây dựng trợ lý AI sử dụng kiến trúc AI Agent hỗ trợ khách hàng ngân hàng\nAI Banking Assistant")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(12.5)
    run_sub.bold = True
    run_sub.font.color.rgb = RGBColor(50, 50, 50)

    # 2. THÔNG TIN METADATA
    meta_box = doc.add_paragraph()
    meta_box.paragraph_format.space_after = Pt(14)
    meta_box.paragraph_format.line_spacing = 1.25

    runs_meta = [
        ("• Ngày báo cáo: ", True), ("19/10/2026 — Tuần từ 12/10/2026 đến 19/10/2026\n", False),
        ("• Giai đoạn: ", True), ("Giai đoạn 1 — Xây dựng cơ sở tri thức và hệ thống RAG\n", False),
        ("• Phiên bản: ", True), ("2.0", False)
    ]
    for text, bold in runs_meta:
        r = meta_box.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        r.bold = bold

    # ĐƯỜNG KẺ NGĂN CÁCH NHẸ
    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(12)
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                     r'</w:pBdr>')
    divider._p.get_or_add_pPr().append(pBdr)

    # 3. PHẦN 1: KẾT QUẢ ĐẠT ĐƯỢC
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Kết quả đạt được trong tuần này — Từ 12/10/2026 đến 19/10/2026")
    r_h1.bold = True
    r_h1.font.name = 'Times New Roman'
    r_h1.font.size = Pt(13)
    r_h1.font.color.rgb = RGBColor(18, 52, 86)

    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(8)
    intro_p.paragraph_format.line_spacing = 1.25
    r_intro = intro_p.add_run(
        "Trong tuần vừa qua, em đã hoàn thành xây dựng phân hệ phân đoạn tài liệu nhận thức bố cục chuyên biệt cho dữ liệu ngân hàng, hỗ trợ toàn diện cả định dạng Word và PDF."
    )
    r_intro.font.name = 'Times New Roman'

    # 1.1 Tóm tắt nguyên lý hoạt động
    h2_1 = doc.add_paragraph()
    h2_1.paragraph_format.space_before = Pt(6)
    h2_1.paragraph_format.space_after = Pt(4)
    r_h2_1 = h2_1.add_run("1.1. Tóm tắt nguyên lý hoạt động của kiến trúc phân đoạn")
    r_h2_1.bold = True
    r_h2_1.font.name = 'Times New Roman'
    r_h2_1.font.size = Pt(12)

    pillars = [
        ("Tiếp nhận và phân loại tài liệu: ", "Hệ thống tự động nhận diện định dạng Word và PDF để chuyển vào các nhánh bóc tách cấu trúc chuyên biệt tương ứng."),
        ("Xử lý bảng biểu phức tạp và chú thích chân trang: ", "Đối với các bảng chứa ô gộp hoặc nhiều tầng tiêu đề, hệ thống chuyển đổi sang dạng hình ảnh để mô hình Vision AI đọc hiểu trọn vẹn quan hệ giữa các hàng và cột. Vùng chú thích dưới đáy bảng được tự động mở rộng và nhúng trực tiếp vào đúng ô dữ liệu mang ký hiệu neo, giúp số liệu không bị tách rời khỏi ngữ cảnh giải thích."),
        ("Xử lý tài liệu scan: ", "Đối với các trang PDF dạng ảnh quét, hệ thống sử dụng công cụ Tesseract OCR v5 kết hợp thư viện xử lý ảnh OpenCV để lọc con dấu mộc đỏ bằng dải màu HSV, cân thẳng góc nghiêng của trang và trích xuất chữ tiếng Việt rõ nét."),
        ("Phân tách câu theo ngữ nghĩa: ", "Thay vì ngắt câu theo quy tắc ký tự thông thường, hệ thống ứng dụng mô hình ngữ nghĩa wtpsplit để nhận diện ranh giới câu tự nhiên theo ngữ cảnh, sau đó gom các câu thành từng đoạn tối ưu từ 150 đến 400 âm tiết mà không bị chồng lấn nội dung."),
        ("Duy trì ngữ cảnh và chuẩn hóa đầu ra: ", "Hệ thống tự động theo dõi cấp độ tiêu đề xuyên suốt tài liệu và gắn kèm tiền tố ngữ cảnh vào đầu mỗi đoạn. Toàn bộ các đoạn văn và bảng biểu sau khi xử lý đều được đóng gói đồng nhất theo cấu trúc 7 trường thông tin, sẵn sàng trực tiếp cho tầng nhúng vector.")
    ]

    for title, desc in pillars:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        r1 = p.add_run(title)
        r1.bold = True
        r1.font.name = 'Times New Roman'
        r2 = p.add_run(desc)
        r2.font.name = 'Times New Roman'

    # 1.2 Sơ đồ kiến trúc
    h2_2 = doc.add_paragraph()
    h2_2.paragraph_format.space_before = Pt(8)
    h2_2.paragraph_format.space_after = Pt(4)
    r_h2_2 = h2_2.add_run("1.2. Sơ đồ kiến trúc phân đoạn tài liệu")
    r_h2_2.bold = True
    r_h2_2.font.name = 'Times New Roman'
    r_h2_2.font.size = Pt(12)

    img_path = r"d:\Desktop\GitHub\AI Banking Asistant\docs\flow diagram\chunking_architecture_simplified.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(img_path, width=Inches(6.2))

        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_p.paragraph_format.space_after = Pt(12)
        r_cap = caption_p.add_run("Hình 1: Sơ đồ kiến trúc và luồng dữ liệu phân đoạn tài liệu nhận thức bố cục")
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(10.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(90, 90, 90)

    # 1.3 Thử nghiệm thực tế
    h2_3 = doc.add_paragraph()
    h2_3.paragraph_format.space_before = Pt(6)
    h2_3.paragraph_format.space_after = Pt(4)
    r_h2_3 = h2_3.add_run("1.3. Kết quả kiểm thử thực nghiệm")
    r_h2_3.bold = True
    r_h2_3.font.name = 'Times New Roman'
    r_h2_3.font.size = Pt(12)

    p_test = doc.add_paragraph()
    p_test.paragraph_format.space_after = Pt(12)
    p_test.paragraph_format.line_spacing = 1.25
    r_test = p_test.add_run(
        "Hệ thống đã được chạy kiểm thử trên toàn bộ tập dữ liệu biểu phí, điều khoản hợp đồng và cẩm nang dịch vụ ngân hàng Ngân hàng. "
        "Kết quả thực nghiệm cho thấy:\n"
        "• Bảo toàn nguyên vẹn số liệu bảng biểu đa tầng và biểu phí phức tạp, không còn tình trạng chú thích chân trang bị mất liên kết với các số liệu tương ứng.\n"
        "• Dữ liệu sạch sẽ, loại bỏ hoàn toàn nhiễu từ con dấu mộc đỏ và mép trang, các đoạn văn bản có độ dài vừa vặn, không bị phân mảnh vụn.\n"
        "• Toàn bộ các đoạn xuất xưởng đều đạt chuẩn 7 trường thông tin, tương thích trực tiếp với các mô hình nhúng ngôn ngữ."
    )
    r_test.font.name = 'Times New Roman'

    # ĐƯỜNG KẺ NGĂN CÁCH NHẸ
    divider2 = doc.add_paragraph()
    divider2.paragraph_format.space_after = Pt(10)
    pBdr2 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                      r'</w:pBdr>')
    divider2._p.get_or_add_pPr().append(pBdr2)

    # 4. PHẦN 2: KẾ HOẠCH TUẦN TỚI
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. Kế hoạch thực hiện tuần tới — Từ 19/10/2026 đến 26/10/2026")
    r_h2.bold = True
    r_h2.font.name = 'Times New Roman'
    r_h2.font.size = Pt(13)
    r_h2.font.color.rgb = RGBColor(18, 52, 86)

    plan_intro = doc.add_paragraph()
    plan_intro.paragraph_format.space_after = Pt(6)
    plan_intro.paragraph_format.line_spacing = 1.25
    r_pi = plan_intro.add_run("Trong tuần tiếp theo, định hướng trọng tâm là đưa các đoạn dữ liệu vào luồng sinh vector và hoàn thiện tầng tìm kiếm:")
    r_pi.font.name = 'Times New Roman'

    plans = [
        ("Triển khai luồng Embedding và cơ sở dữ liệu Vector: ", "Tích hợp các đoạn dữ liệu vào mô hình nhúng ngôn ngữ và thiết lập kho lưu trữ vector ngữ nghĩa song song với chỉ mục tìm kiếm từ khóa."),
        ("Phát triển tầng Retrieval và xếp hạng lại: ", "Xây dựng cơ chế tìm kiếm lai kết hợp giữa tìm kiếm từ khóa và tìm kiếm ngữ nghĩa, tích hợp thêm mô hình xếp hạng lại để tinh chọn các đoạn câu trả lời chính xác nhất."),
        ("Tiếp tục tinh chỉnh phân đoạn: ", "Đánh giá chất lượng tìm kiếm trên tập câu hỏi ngân hàng thực tế, tinh chỉnh độ dài đoạn văn và tối ưu hóa thời gian xử lý bảng biểu phức tạp.")
    ]

    for title, desc in plans:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        r1 = p.add_run(title)
        r1.bold = True
        r1.font.name = 'Times New Roman'
        r2 = p.add_run(desc)
        r2.font.name = 'Times New Roman'

    # LƯU FILE
    out_dir = r"d:\Desktop\GitHub\AI Banking Asistant\docs\báo cáo tiến độ"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Báo cáo tiến độ v2.docx")
    alt_path = os.path.join(out_dir, "Báo cáo tiến độ v2_moi.docx")
    try:
        doc.save(out_path)
        print("Report DOCX generated successfully at:", out_path)
    except PermissionError:
        doc.save(alt_path)
        print("Note: 'Báo cáo tiến độ v2.docx' đang được mở trong Word.")
        print("Bản mới đã được lưu thành công tại:", alt_path)

if __name__ == "__main__":
    create_report()
