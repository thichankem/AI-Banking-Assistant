# -*- coding: utf-8 -*-
"""
Kịch bản tự động tạo Báo cáo Tiến độ v3 (Bản Word .docx) - Phiên bản Chuẩn hóa Thuần DOCX
Đề tài: Xây dựng trợ lý AI sử dụng kiến trúc AI Agent hỗ trợ khách hàng ngân hàng (AI Banking Assistant)
Thời gian: Tuần từ 19/10/2026 đến 26/10/2026 (Ngày lập: 24/10/2026)
Quy chuẩn:
- 100% loại bỏ toàn bộ nội dung PDF (Hệ thống chuẩn hóa hoàn toàn trên định dạng văn bản ngân hàng .docx)
- Phong cách học thuật, trang trọng (xưng "em", không dùng "nhóm nghiên cứu")
- 100% chữ màu đen / tương phản cao, bảng có chia cột lưới rõ ràng
- Tập trung vào: Kiến trúc 7 Node RAG Folder-Driven, Bộ bóc tách DOCX chuyên sâu, Web UI & API Server,
  Kiểm toán Chaos Engineering (43/43 Pass), và Kế hoạch triển khai Hệ thống Bộ nhớ (Memory Storage) tuần tới.
"""

import os
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_MUTED = RGBColor(90, 90, 90)
COLOR_PRIMARY = RGBColor(18, 52, 86)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="777777", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def create_grid_table(doc, headers, data, widths, aligns=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "EEEEEE")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_borders(hdr_cells[i], top="single", bottom="single", left="single", right="single", color="444444", sz="6")
        
        p = hdr_cells[i].paragraphs[0]
        if aligns and i < len(aligns):
            p.alignment = aligns[i]
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = COLOR_BLACK

    for row_idx, row_val in enumerate(data):
        row = tbl.rows[row_idx + 1]
        bg_color = "F9F9F9" if (row_idx % 2 == 1) else "FFFFFF"

        for i, val in enumerate(row_val):
            cell = row.cells[i]
            cell.text = str(val)
            cell.width = widths[i]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=130, right=130)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="999999", sz="4")
            
            p = cell.paragraphs[0]
            if aligns and i < len(aligns):
                p.alignment = aligns[i]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_BLACK
                    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)
    return tbl

def add_p(doc, text="", bold_prefix="", italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.italic = italic
        r.font.color.rgb = COLOR_BLACK
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_BLACK
    return p

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK
    return h

def generate_report():
    doc = docx.Document()

    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    # 1. TIÊU ĐỀ BÁO CÁO
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(3)
    
    r_t = p_title.add_run("BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI GR2\n")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_PRIMARY
    
    r_sub = p_title.add_run("Đề tài: Xây dựng trợ lý AI sử dụng kiến trúc AI Agent hỗ trợ khách hàng ngân hàng\n")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_BLACK
    
    r_eng = p_title.add_run("AI Banking Assistant")
    r_eng.font.name = "Times New Roman"
    r_eng.font.size = Pt(11.5)
    r_eng.font.italic = True
    r_eng.font.color.rgb = COLOR_MUTED

    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_meta.paragraph_format.space_before = Pt(10)
    p_meta.paragraph_format.space_after = Pt(8)
    p_meta.paragraph_format.line_spacing = 1.2
    
    r_m = p_meta.add_run("• Ngày báo cáo: ")
    r_m.bold = True
    p_meta.add_run("24/10/2026 (Tuần từ 19/10/2026 đến 26/10/2026)\n")
    
    r_m = p_meta.add_run("• Giai đoạn: ")
    r_m.bold = True
    p_meta.add_run("Giai đoạn 1 — Hoàn thiện Kiến trúc RAG Đa tầng & Khởi động Tầng Bộ nhớ Agentic\n")
    
    r_m = p_meta.add_run("• Phiên bản: ")
    r_m.bold = True
    p_meta.add_run("3.0 (Bản chuẩn hóa Ingestion tài liệu nghiệp vụ)")

    # Divider
    divider = doc.add_paragraph()
    divider.paragraph_format.space_after = Pt(8)
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                     r'</w:pBdr>')
    divider._p.get_or_add_pPr().append(pBdr)

    # =========================================================================
    # PHẦN 1: KẾT QUẢ ĐẠT ĐƯỢC
    # =========================================================================
    add_heading_1(doc, "1. Kết quả đạt được trong tuần — Từ 19/10/2026 đến 26/10/2026")
    add_p(doc, "Trong tuần vừa qua, em đã hoàn thành việc nâng cấp toàn diện hệ thống từ phân đoạn đơn lẻ thành một đường ống RAG chuyên sâu 7 Node hoàn chỉnh, chuẩn hóa 100% kho tài liệu tri thức ngân hàng theo định dạng Microsoft Word (.docx), triển khai máy chủ ứng dụng Web tương tác và thiết lập bộ kiểm định chất lượng thực nghiệm đạt tỷ lệ 100% Pass tuyệt đối.")

    add_heading_2(doc, "1.1. Chuẩn hóa 100% Kho Tri thức Ngân hàng & Nâng cấp DOCX Chunker")
    add_p(doc, "Để bảo đảm tính toàn vẹn cao nhất của các biểu phí, điều khoản hợp đồng và cẩm nang ngân hàng, toàn bộ kho dữ liệu được chuẩn hóa thống nhất theo định dạng Microsoft Word (.docx). Phân hệ DOCX Chunker được nâng cấp chuyên sâu với các tiêu chuẩn kỹ thuật vượt trội:")
    add_bullet(doc, "Bảo toàn hoàn toàn cấu trúc chữ in đậm (**text**), in nghiêng (*text*) và gạch chân (<u>text</u>) trực tiếp từ thuộc tính định dạng Word Run, giúp bảo toàn tính trang trọng và ngữ nghĩa của các điều khoản pháp lý.", bold_prefix="Bảo tồn định dạng ngữ nghĩa: ")
    add_bullet(doc, "Tự động gộp các Run liền kề cùng định dạng để tránh hiện tượng vỡ từ ngữ; sửa triệt để lỗi dính chữ sau nhãn Markdown (ví dụ: '**Ghi chú:**Đối với' thành '**Ghi chú:** Đối với') và sau dấu đóng ngoặc đơn kết hợp chữ in hoa.", bold_prefix="Chuẩn hóa khoảng trắng và liên kết từ: ")
    add_bullet(doc, "Ứng dụng phân tích cây OpenXML AST để bóc tách hình ảnh, bảng biểu minh họa, áp dụng điều kiện tiên quyết: chỉ lưu hình ảnh khi có chú thích (caption) bên dưới, loại bỏ hoàn toàn các logo, icon và chi tiết trang trí gây nhiễu.", bold_prefix="Trích xuất biểu đồ có chú thích: ")
    add_bullet(doc, "Phân đoạn câu theo mô hình wtpsplit và kiểm soát ngân sách âm tiết từ 200 đến 380 âm tiết (trần cứng 420 âm tiết), bảo toàn đường dẫn phân cấp Breadcrumbs (Tên tài liệu > Mục lớn > Khoản mục).", bold_prefix="Kiểm soát ngân sách âm tiết co giãn: ")

    add_heading_2(doc, "1.2. Hoàn thiện Đường ống Dữ liệu RAG 7 Node Liên hoàn (Folder-to-Folder)")
    add_p(doc, "Hệ thống được thiết kế theo kiến trúc Folder-to-Folder minh bạch, trong đó mỗi Node đều có thư mục Đầu vào (Input) và Đầu ra (Output) dạng JSON để đối soát độc lập:")

    nodes_info = [
        ("Node 1 — Phân đoạn tài liệu (Chunking Node): ", "Đọc tài liệu Word chuẩn hóa và xuất các đoạn trích đạt chuẩn cấu trúc 7 trường thông tin (chunk_id, source, breadcrumb, section_title, content, original_text, parent_id)."),
        ("Node 2 — Nhúng vector và lưu bộ đệm (Embedding Node): ", "Tạo vector 1024 chiều bằng mô hình multilingual-e5-large (FPT Cloud API) kết hợp cơ chế lưu bộ đệm SQLite Cache đa luồng (WAL mode), loại bỏ 100% thời gian gọi lại API với nội dung trùng lặp."),
        ("Node 3 — Chỉ mục kép Qdrant & BM25 (Indexing Node): ", "Lập chỉ mục song song vào Qdrant HNSW 1024-dim và bảng chỉ mục ngược Okapi BM25, tích hợp cơ chế tự động chuyển sang In-Memory khi xảy ra xung đột khóa tệp."),
        ("Node 4 — Chuẩn hóa câu hỏi (Query Parser Node): ", "Tự động nhận diện và mở rộng từ viết tắt ngân hàng (eKYC, TKTT, Smart OTP, Digibank) và trích xuất chuẩn xác các tham chiếu điều khoản (Điều X, Khoản Y)."),
        ("Node 5 — Tìm kiếm lai kết hợp (Hybrid Retrieval Node): ", "Thực thi truy vấn đồng thời trên Dense Vector và BM25, sau đó hợp nhất bằng thuật toán Reciprocal Rank Fusion (RRF k=60) để lấy Top-5 ứng viên cân bằng cả về ngữ nghĩa lẫn từ khóa chính xác."),
        ("Node 6 — Tái xếp hạng ngữ nghĩa (Cross-Encoder Reranker Node): ", "Sử dụng mô hình bge-reranker-v2-m3 để chấm điểm chéo cặp câu hỏi - đoạn trích, kéo các đoạn văn bản trả lời đúng nhất lên Top 3."),
        ("Node 7 — Sinh câu trả lời & Phòng vệ Guardrails Node): ", "Mô hình DeepSeek-V4-Flash sinh câu trả lời trích dẫn nguồn cụ thể. Hệ thống Banking Guardrails bảo vệ hai chiều: chặn đứng Prompt Injection (kể cả dạng mã hóa Leetspeak), che giấu PII (CCCD, tài khoản, số thẻ tín dụng 16 số), và cảnh báo lừa đảo chiếm đoạt OTP/mật khẩu.")
    ]
    for n_title, n_desc in nodes_info:
        add_bullet(doc, n_desc, bold_prefix=n_title)

    add_heading_2(doc, "1.3. Xây dựng Máy chủ Web API và Giao diện Trực quan hóa")
    add_bullet(doc, "Triển khai máy chủ FastAPI (main.py) cung cấp đầy đủ các cổng giao tiếp API (/api/chat, /api/config, /api/nodes, /api/pipeline/run-folders), phục vụ trực tiếp giao diện Web tại địa chỉ http://127.0.0.1:8000.", bold_prefix="Máy chủ API chuyên dụng: ")
    add_bullet(doc, "Giao diện Web tương tác (HTML/CSS/JS thuần) tích hợp khung trò chuyện thời gian thực, hiển thị trích dẫn nguồn điều khoản, thẻ phân loại ý định và thời gian phản hồi.", bold_prefix="Giao diện Khách hàng: ")
    add_bullet(doc, "Cho phép duyệt cây thư mục dữ liệu của cả 7 Node, xem trước nội dung tệp JSON đầu vào/đầu ra, giúp quá trình đối soát và nghiệm thu đề tài minh bạch 100%.", bold_prefix="Bảng điều khiển Node Explorer: ")

    add_heading_2(doc, "1.4. Kiểm toán Chất lượng Thực nghiệm & Khắc phục Lỗ hổng (Chaos Engineering)")
    add_p(doc, "Để đảm bảo hệ thống đạt độ tin cậy cấp độ ngân hàng, em đã xây dựng bộ kiểm định chất lượng thực nghiệm gồm 43 kịch bản cực hạn (Chaos Engineering) chia thành 5 cấp độ:")

    bench_headers = ["Chỉ số Kiểm định (Metric)", "Kết quả Đạt được", "Đánh giá Kỹ thuật"]
    bench_data = [
        ["Độ bao phủ dữ liệu tài liệu Word", "100.0% (56/56 tệp)", "Chuẩn hóa hoàn toàn kho tri thức DOCX"],
        ["Tính nguyên vẹn ranh giới câu", "100.0% Pass", "Áp dụng bộ ngắt câu ngữ nghĩa wtpsplit"],
        ["Bảo tồn Bảng biểu & Số liệu", "100.0% Pass", "Giữ nguyên cấu trúc ma trận bảng biểu phí"],
        ["Hit Rate@1 (Top 1 đúng nguồn)", "83.33% (15/18 truy vấn)", "RRF và Dense định vị chính xác tài liệu"],
        ["Hit Rate@3 (Top 3 đúng nguồn)", "88.89% (16/18 truy vấn)", "Cung cấp đầy đủ ngữ cảnh cho LLM tổng hợp"],
        ["Hit Rate@5 (Top 5 đúng nguồn)", "94.44% (17/18 truy vấn)", "Tỷ lệ thu hút tài liệu liên quan xuất sắc"],
        ["Mean Reciprocal Rank (MRR)", "0.8722", "Điểm số xếp hạng tương đối cao"],
        ["Kháng bẫy Hallucination", "100.0% (2/2 bẫy bị chặn)", "Bác bỏ gói vay 0% và lãi suất 18% bịa đặt"],
        ["Phòng vệ Guardrails (OTP / PII)", "100.0% Pass", "Chặn đứng Prompt Injection Leetspeak & che số thẻ"],
        ["Bộ kiểm thử Hỗn loạn (Chaos Test)", "43/43 Ca kiểm thử Vượt qua (100%)", "Vá xong toàn bộ 3 lỗi và lỗ hổng thực tế"]
    ]
    bench_widths = [Inches(2.5), Inches(1.8), Inches(2.4)]
    create_grid_table(doc, bench_headers, bench_data, bench_widths,
                      aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])

    add_p(doc, "Quá trình kiểm toán đã bóc tách và sửa chữa triệt để 3 lỗi/lỗ hổng thực tế: (1) Lỗi AttributeError khi chunk rỗng trong Embedder; (2) Lỗ hổng để lọt số thẻ tín dụng phân cách bằng dấu cách (4532 0156 7890 1234) trong Guardrail; và (3) Lỗ hổng vượt rào bằng ký tự Leetspeak (1gnore...). Toàn bộ 43 ca kiểm thử hiện nay đều đạt 100% Pass trong 14.09 giây.")

    # Divider
    divider2 = doc.add_paragraph()
    divider2.paragraph_format.space_after = Pt(8)
    pBdr2 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                      r'</w:pBdr>')
    divider2._p.get_or_add_pPr().append(pBdr2)

    # =========================================================================
    # PHẦN 2: KẾ HOẠCH TUẦN TỚI
    # =========================================================================
    add_heading_1(doc, "2. Kế hoạch thực hiện tuần tới — Từ 26/10/2026 đến 02/11/2026")
    add_p(doc, "Sau khi hoàn thiện vững chắc nền tảng RAG đa tầng, trọng tâm nghiên cứu và triển khai trong tuần tiếp theo là nâng cấp hệ thống từ mô hình hỏi đáp đơn lượt (Stateless RAG) sang Kiến trúc Trợ lý AI có Trạng thái (Stateful AI Agent), bắt đầu triển khai toàn diện Phân hệ Lưu trữ Bộ nhớ (Memory Storage) cho bot:")

    plans = [
        ("Xây dựng Kiến trúc Bộ nhớ Đa tầng (Multi-Tier Memory Architecture): ",
         "Thiết lập cơ chế bộ nhớ 3 cấp: (1) Bộ nhớ làm việc ngắn hạn (Working/Short-term Memory) lưu trữ ngữ cảnh hội thoại trượt (Sliding Window Buffer) giúp bot duy trì mạch đàm thoại liên tục; "
         "(2) Bộ nhớ dài hạn (Long-term Episodic & Semantic Memory) lưu trữ thông tin sở thích, lịch sử giao dịch và hồ sơ định danh của khách hàng trên cơ sở dữ liệu SQLite; "
         "(3) Bộ nhớ tóm tắt tự động (Hierarchical Conversation Summarization) định kỳ nén các phiên đàm thoại dài bằng LLM để giữ trọn vẹn ngữ cảnh mà không làm tràn cửa sổ ngữ cảnh."),
        ("Phát triển Phân hệ Viết lại Truy vấn theo Ngữ cảnh (Contextual Query Rewriter): ",
         "Xây dựng Node xử lý ngôn ngữ trung gian đứng trước tầng Retrieval, có nhiệm vụ phân tích lịch sử hội thoại từ bộ nhớ để tự động bổ sung chủ ngữ, đại từ khuyết thiếu trong câu hỏi của khách hàng (ví dụ: khi khách hàng hỏi 'Phí của nó là bao nhiêu?', hệ thống tự động suy luận 'nó' là 'Thẻ Vietcombank Visa Signature' đã đề cập ở câu trước để tìm kiếm chính xác)."),
        ("Tích hợp Trình quản lý Phiên (Session Manager & Checkpointer): ",
         "Ứng dụng LangGraph Checkpointer hoặc thiết kế Session Store chuyên biệt trên SQLite, cho phép lưu trữ vết trạng thái đa phiên, hỗ trợ khách hàng tiếp tục cuộc trò chuyện sau khi tải lại trang web hoặc chuyển đổi thiết bị."),
        ("Nâng cấp Giao diện Web hỗ trợ Quản lý Phiên hội thoại: ",
         "Bổ sung thanh điều hướng danh sách các phiên trò chuyện (Chat History Sidebar) ở giao diện Frontend, cho phép tạo phiên hội thoại mới, đổi tên phiên và xem lại lịch sử tư vấn tài chính theo thời gian thực.")
    ]

    for p_title, p_desc in plans:
        add_bullet(doc, p_desc, bold_prefix=p_title)

    # Lưu tệp Word
    out_dir = r"d:\Desktop\GitHub\AI Banking Asistant\docs\báo cáo tiến độ"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Báo cáo tiến độ v3.docx")
    alt_path = os.path.join(out_dir, "Báo cáo tiến độ v3_moi.docx")

    try:
        doc.save(out_path)
        print("Report DOCX generated successfully at:", out_path)
    except PermissionError:
        doc.save(alt_path)
        print("Note: 'Báo cáo tiến độ v3.docx' đang mở. Đã lưu tại:", alt_path)

if __name__ == "__main__":
    generate_report()
