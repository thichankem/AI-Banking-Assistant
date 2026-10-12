"""
Excel Table Layout-Aware Chunker
Trích xuất và phân đoạn các bảng biểu ma trận tài chính phức tạp (.xlsx) của Vietcombank:
- Đọc đa sheet (multi-sheet) qua openpyxl.
- Tự động nhận diện dòng tiêu đề cột (header row).
- Làm phẳng từng hàng dữ liệu kèm đầy đủ tên cột (Linearized Table Rows),
  giúp mô hình vector và tìm kiếm không bao giờ mất ngữ cảnh cột khi tra cứu số liệu.
"""

from pathlib import Path
import openpyxl

class ExcelChunker:
    def __init__(self, rows_per_chunk=3):
        self.rows_per_chunk = rows_per_chunk

    def chunk_excel(self, excel_path):
        """
        Đọc và phân đoạn bảng biểu từ file Excel .xlsx
        """
        excel_path = Path(excel_path)
        chunks = []

        try:
            wb = openpyxl.load_workbook(str(excel_path), data_only=True)
        except Exception as e:
            print(f"[Excel Chunker] Không thể mở file {excel_path.name}: {e}")
            return []

        chunk_counter = 0

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                continue

            # 1. Tìm dòng header đầu tiên có dữ liệu
            header_idx = -1
            headers = []
            for r_idx, row in enumerate(rows):
                non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if len(non_empty) >= 2:
                    header_idx = r_idx
                    headers = [str(c).strip() if c is not None else f"Cột_{i+1}" for i, c in enumerate(row)]
                    break

            if header_idx == -1:
                continue

            # 2. Xử lý các dòng dữ liệu sau header
            data_rows = rows[header_idx + 1:]
            buffer_row_strs = []
            buffer_raw = []

            breadcrumb = f"{excel_path.stem} > Sheet: {sheet_name}"

            for row_idx, row in enumerate(data_rows, start=header_idx + 2):
                if not any(row):
                    continue

                # Ghép từng cell với tên cột tương ứng
                row_items = []
                for c_idx, cell_val in enumerate(row):
                    if cell_val is None or str(cell_val).strip() == "":
                        continue
                    col_name = headers[c_idx] if c_idx < len(headers) else f"Cột_{c_idx+1}"
                    row_items.append(f"{col_name}: {cell_val}")

                if not row_items:
                    continue

                row_text = " | ".join(row_items)
                buffer_row_strs.append(f"* Hàng {row_idx}: {row_text}")
                buffer_raw.append(row_text)

                # Khi buffer đạt số lượng dòng yêu cầu, đóng gói thành 1 chunk
                if len(buffer_row_strs) >= self.rows_per_chunk:
                    full_content = (
                        f"[{breadcrumb} > Bảng dữ liệu]\n"
                        f"Cột dữ liệu: {', '.join([h for h in headers if h])}\n"
                        + "\n".join(buffer_row_strs)
                    )
                    chunks.append({
                        "chunk_id": f"{excel_path.stem}_{sheet_name}_c{chunk_counter}",
                        "source": excel_path.name,
                        "breadcrumb": f"{breadcrumb} > Bảng dữ liệu",
                        "section_title": f"Bảng {sheet_name}",
                        "chunk_type": "excel_table",
                        "content": full_content,
                        "raw_content": "\n".join(buffer_raw)
                    })
                    chunk_counter += 1
                    buffer_row_strs = []
                    buffer_raw = []

            # Lưu các dòng còn lại trong buffer
            if buffer_row_strs:
                full_content = (
                    f"[{breadcrumb} > Bảng dữ liệu]\n"
                    f"Cột dữ liệu: {', '.join([h for h in headers if h])}\n"
                    + "\n".join(buffer_row_strs)
                )
                chunks.append({
                    "chunk_id": f"{excel_path.stem}_{sheet_name}_c{chunk_counter}",
                    "source": excel_path.name,
                    "breadcrumb": f"{breadcrumb} > Bảng dữ liệu",
                    "section_title": f"Bảng {sheet_name}",
                    "chunk_type": "excel_table",
                    "content": full_content,
                    "raw_content": "\n".join(buffer_raw)
                })
                chunk_counter += 1

        return chunks
