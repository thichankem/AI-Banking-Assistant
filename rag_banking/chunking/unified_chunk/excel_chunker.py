"""
Module Excel Table Layout-Aware Chunker
========================================
Bộ trích xuất và phân đoạn bảng biểu ma trận tài chính (.xlsx) ngân hàng:
1. Multi-Sheet Extraction: Quét toàn diện các sheet trong file Excel qua openpyxl.
2. Dynamic Header Row Detection: Tự động phát hiện hàng tiêu đề thực tế của bảng dữ liệu.
3. Linearized Table Rows: Tuyến tính hóa từng hàng dạng "Cột: Giá trị", bảo toàn quan hệ ngữ cảnh cột cho Vector Embeddings & BM25.
4. Batch Row Chunking: Nhóm theo số hàng hợp lý (mặc định 3 hàng) thành 1 chunk kèm breadcrumb rõ ràng.
5. Schema Compliance: Chuẩn hóa Schema tinh gọn: có 'source', KHÔNG CÓ 'ref', 'legal_ref', 'raw_content', 'char_length', 'format'.
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Any, List, Optional, Union
import openpyxl


class ExcelChunker:
    """
    Bộ phân đoạn tài liệu bảng biểu Excel chuyên dụng cho ngân hàng theo giải thuật DLA.
    """

    def __init__(self, rows_per_chunk: int = 3):
        self.rows_per_chunk = rows_per_chunk

    def chunk_excel(self, excel_path: Union[str, Path]) -> List[Dict[str, Any]]:
        """
        Đọc và phân đoạn các bảng biểu ma trận từ tệp Excel .xlsx.
        """
        excel_path = Path(excel_path)
        if not excel_path.exists():
            print(f"[Excel Chunker] Không tìm thấy tệp: {excel_path}")
            return []

        try:
            wb = openpyxl.load_workbook(str(excel_path), data_only=True)
        except Exception as e:
            print(f"[Excel Chunker] Không thể mở tệp {excel_path.name}: {e}")
            return []

        chunks: List[Dict[str, Any]] = []
        chunk_counter = 0

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                continue

            # 1. Tìm hàng header đầu tiên có ít nhất 2 cột dữ liệu hợp lệ
            header_idx = -1
            headers: List[str] = []
            for r_idx, row in enumerate(rows):
                non_empty = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if len(non_empty) >= 2:
                    header_idx = r_idx
                    headers = [
                        str(c).strip() if c is not None and str(c).strip() else f"Cột_{i+1}"
                        for i, c in enumerate(row)
                    ]
                    break

            if header_idx == -1:
                continue

            # 2. Tuyến tính hóa và nhóm dữ liệu sau hàng header
            data_rows = rows[header_idx + 1:]
            buffer_row_strs: List[str] = []

            breadcrumb = f"{excel_path.stem} > Sheet: {sheet_name} > Bảng dữ liệu"
            section_title = f"Bảng {sheet_name}"

            for row_idx, row in enumerate(data_rows, start=header_idx + 2):
                if not any(row):
                    continue

                # Ghép từng ô với tên cột tương ứng
                row_items = []
                for c_idx, cell_val in enumerate(row):
                    if cell_val is None or str(cell_val).strip() == "":
                        continue
                    col_name = headers[c_idx] if c_idx < len(headers) else f"Cột_{c_idx+1}"
                    row_items.append(f"{col_name}: {cell_val}")

                if not row_items:
                    continue

                row_text = " | ".join(row_items)
                buffer_row_strs.append(f"* Hàng {row_idx}: {row_text}")

                # Khi buffer đạt rows_per_chunk, đóng gói thành 1 chunk
                if len(buffer_row_strs) >= self.rows_per_chunk:
                    chunk_obj = self._create_excel_chunk(
                        excel_path=excel_path,
                        sheet_name=sheet_name,
                        chunk_counter=chunk_counter,
                        breadcrumb=breadcrumb,
                        section_title=section_title,
                        headers=headers,
                        buffer_row_strs=buffer_row_strs
                    )
                    chunks.append(chunk_obj)
                    chunk_counter += 1
                    buffer_row_strs.clear()

            # Đóng gói các dòng còn lại trong buffer
            if buffer_row_strs:
                chunk_obj = self._create_excel_chunk(
                    excel_path=excel_path,
                    sheet_name=sheet_name,
                    chunk_counter=chunk_counter,
                    breadcrumb=breadcrumb,
                    section_title=section_title,
                    headers=headers,
                    buffer_row_strs=buffer_row_strs
                )
                chunks.append(chunk_obj)
                chunk_counter += 1

        return chunks

    def _create_excel_chunk(
        self,
        excel_path: Path,
        sheet_name: str,
        chunk_counter: int,
        breadcrumb: str,
        section_title: str,
        headers: List[str],
        buffer_row_strs: List[str]
    ) -> Dict[str, Any]:
        """Đóng gói các dòng buffer thành một đối tượng chunk Excel chuẩn Schema tinh gọn."""
        full_content = (
            f"[Breadcrumbs: {breadcrumb}]\n"
            f"Cột dữ liệu: {', '.join([h for h in headers if h])}\n"
            f"Chi tiết từng hàng:\n" + "\n".join(buffer_row_strs)
        )

        return {
            "chunk_id": f"{excel_path.stem}_{sheet_name}_c{chunk_counter}",
            "source": excel_path.name,
            "chunk_type": "excel_table",
            "breadcrumb": breadcrumb,
            "section_title": section_title,
            "sheet_name": sheet_name,
            "content": full_content,
            "word_count": len(full_content.split())
        }
