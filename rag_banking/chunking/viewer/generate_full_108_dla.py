#!/usr/bin/env python3
"""
Bộ trích xuất và tiền xử lý Layout DLA & Chunks chuẩn 108 tệp Vietcombank
======================================================================
TUYỆT ĐỐI 0% REGEX (Zero Regex-Based Hierarchy):
- Đối với PDF: Sử dụng mô hình hình học không gian 5 chiều OCRLayoutHierarchyPredictor
  (rel_h, gap_before, is_centered, left_indent, rel_w) kết hợp PyMuPDF find_tables.
- Đối với DOCX: Sử dụng Character-Weighted Font Size Histogram và OpenXML AST styles.
- Tự động liên kết 2 chiều giữa Visual Bounding Boxes và Chunks.
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path
from collections import Counter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

# Đảm bảo import được các module của dự án
CURRENT_DIR = Path(__file__).resolve().parent
CHUNKING_ROOT = CURRENT_DIR.parent
PROJECT_ROOT = CHUNKING_ROOT.parent.parent

sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(CHUNKING_ROOT))

import pymupdf
import docx
from rag_banking.chunking.pdf_chunker.ocr_layout_predictor import (
    OCRLayoutHierarchyPredictor,
    LayoutHierarchyLevel
)
from rag_banking.chunking.pdf_chunker.pdf_chunker import (
    PDFChunker,
    extract_document_title_dla,
    clean_filename_fallback,
    is_valid_title_candidate
)

PDF_DIR = PROJECT_ROOT / "dataset" / "chunking_dataset" / "pdf"
DOCX_DIR = PROJECT_ROOT / "dataset" / "chunking_dataset" / "docx"
CHUNK_CANDIDATE_PATHS = [
    CHUNKING_ROOT / "chunking_dataset" / "output_dataset" / "all_chunks.json",
    PROJECT_ROOT / "dataset" / "chunking_dataset" / "output_dataset" / "all_chunks.json",
    CURRENT_DIR / "sample_data" / "all_chunks_complete.json",
]
OUTPUT_LAYOUTS_DIR = CURRENT_DIR / "layouts"
OUTPUT_MANIFEST_JSON = CURRENT_DIR / "manifest_108.json"
OUTPUT_MANIFEST_JS = CURRENT_DIR / "manifest_108.js"

# Tạo thư mục đầu ra nếu chưa có
OUTPUT_LAYOUTS_DIR.mkdir(parents=True, exist_ok=True)

# Khởi tạo bộ dự đoán DLA thuần hình học và bộ phân đoạn PDF
dla_predictor = OCRLayoutHierarchyPredictor()
pdf_chunker_instance = PDFChunker()


def load_all_chunks() -> dict[str, list[dict]]:
    """Tải và phân nhóm toàn bộ chunks đã xử lý theo tên tệp nguồn."""
    chosen_path = None
    for p in CHUNK_CANDIDATE_PATHS:
        if p.exists():
            chosen_path = p
            break

    if not chosen_path:
        print(f"[!] Không tìm thấy tệp chunks nào trong các đường dẫn chỉ định.")
        return {}

    print(f"[*] Sử dụng tệp Chunks: {chosen_path}")
    with open(chosen_path, "r", encoding="utf-8") as f:
        chunks_list = json.load(f)

    grouped: dict[str, list[dict]] = {}
    for c in chunks_list:
        src = c.get("source", "").strip()
        if src:
            grouped.setdefault(src, []).append(c)
    return grouped


def process_pdf_document(
    pdf_path: Path,
    chunks: list[dict]
) -> dict:
    """
    Trích xuất layout elements cho tệp PDF bằng PyMuPDF và 5D DLA Geometry Predictor (0% Regex).
    """
    doc = pymupdf.open(pdf_path)
    total_pages = len(doc)
    doc_name = pdf_path.name

    pages_data = []
    all_elements = []
    chunk_to_elements = {}

    # Giới hạn xử lý tối đa 15 trang đầu cho các tài liệu quá dài để tối ưu hiệu năng
    max_process_pages = min(total_pages, 15)

    for p_idx in range(max_process_pages):
        page = doc[p_idx]
        p_num = p_idx + 1
        page_w = float(page.rect.width)
        page_h = float(page.rect.height)

        # 1. Phát hiện bảng biểu bằng PyMuPDF table detection
        table_bboxes = []
        try:
            tabs = page.find_tables()
            for t in tabs.tables:
                table_bboxes.append(t.bbox)
        except Exception:
            pass

        # 2. Trích xuất text blocks và lines
        text_dict = page.get_text("dict")
        raw_lines = []

        for b_idx, block in enumerate(text_dict.get("blocks", [])):
            if "lines" not in block:
                continue
            for l in block["lines"]:
                line_text = "".join(s["text"] for s in l["spans"]).strip()
                if not line_text:
                    continue
                x0, y0, x1, y1 = l["bbox"]
                w = max(x1 - x0, 1.0)
                h = max(y1 - y0, 1.0)
                raw_lines.append({
                    "text": line_text,
                    "x_min": x0,
                    "y_min": y0,
                    "x_max": x1,
                    "y_max": y1,
                    "width": w,
                    "height": h,
                    "center_x": (x0 + x1) / 2.0
                })

        # Nếu là trang scan/slide ảnh thuần (không có digital text), trích xuất OCR lines thực tế kèm bounding box chính xác
        if not raw_lines and pdf_chunker_instance.ocr_extractor.is_available():
            try:
                img_np = pdf_chunker_instance.ocr_extractor.render_pdf_page_300dpi(pdf_path, page_idx=p_idx)
                if img_np is not None:
                    ocr_lines = pdf_chunker_instance.ocr_extractor.extract_layout_lines(img_np)
                    if ocr_lines:
                        img_h, img_w = img_np.shape[:2]
                        scale_x = page_w / float(img_w)
                        scale_y = page_h / float(img_h)
                        for ol in ocr_lines:
                            txt = ol.get("text", "").strip()
                            if not txt:
                                continue
                            x0 = ol.get("x_min", 0.0) * scale_x
                            y0 = ol.get("y_min", 0.0) * scale_y
                            x1 = ol.get("x_max", 0.0) * scale_x
                            y1 = ol.get("y_max", 0.0) * scale_y
                            raw_lines.append({
                                "text": txt,
                                "x_min": x0,
                                "y_min": y0,
                                "x_max": x1,
                                "y_max": y1,
                                "width": max(x1 - x0, 1.0),
                                "height": max(y1 - y0, 1.0),
                                "center_x": (x0 + x1) / 2.0
                            })
            except Exception:
                pass

        # 3. Phân loại cấp bậc bằng OCRLayoutHierarchyPredictor (0% Regex)
        page_elements = []
        if raw_lines:
            stats = dla_predictor.compute_page_layout_stats(raw_lines, page_w, page_h)
            classified = dla_predictor.classify_page_lines(raw_lines, page_w, page_h, page_num=p_num)

            for idx, item in enumerate(classified):
                eid = f"{doc_name}_p{p_num}_e{idx}"
                bbox = item.get("box", [item.get("x_min", 0), item.get("y_min", 0), item.get("x_max", 0), item.get("y_max", 0)])
                
                # Kiểm tra nếu nằm trong vùng bảng
                is_in_table = any(
                    (bbox[0] >= tb[0] - 5 and bbox[1] >= tb[1] - 5 and bbox[2] <= tb[2] + 5 and bbox[3] <= tb[3] + 5)
                    for tb in table_bboxes
                )

                level = LayoutHierarchyLevel.BODY
                if is_in_table:
                    level = "table"
                else:
                    level = item.get("predicted_level", LayoutHierarchyLevel.BODY)

                el_entry = {
                    "element_id": eid,
                    "page_num": p_num,
                    "box": [round(float(coord), 1) for coord in bbox],
                    "text": item.get("text", ""),
                    "type": level,
                    "confidence": item.get("confidence", 0.95),
                    "chunk_id": None
                }
                page_elements.append(el_entry)
                all_elements.append(el_entry)
        else:
            # Nếu là trang scan ảnh thuần, trích xuất text từ chunks đã OCR của trang đó
            matching_chunks = [
                c for c in chunks
                if p_num in (c.get("page_numbers") or [c.get("page_start", 0)])
            ]
            if matching_chunks:
                y_cursor = page_h * 0.10
                for c_idx, c in enumerate(matching_chunks):
                    lines = [ln.strip() for ln in c.get("content", "").split("\n") if ln.strip()]
                    for l_idx, line in enumerate(lines[:10]):
                        eid = f"{doc_name}_p{p_num}_scan_e{c_idx}_{l_idx}"
                        lh = 18.0
                        el_entry = {
                            "element_id": eid,
                            "page_num": p_num,
                            "box": [round(50.0, 1), round(y_cursor, 1), round(page_w - 50.0, 1), round(y_cursor + lh, 1)],
                            "text": line[:120],
                            "type": "title" if (l_idx == 0 and len(line) < 60) else "body",
                            "confidence": 0.90,
                            "chunk_id": c.get("chunk_id")
                        }
                        y_cursor += lh + 6.0
                        page_elements.append(el_entry)
                        all_elements.append(el_entry)

        pages_data.append({
            "page_num": p_num,
            "width": page_w,
            "height": page_h,
            "elements": page_elements
        })

    # 4. Nếu là slide (7 file slide) chưa có trong all_chunks.json, tự động tạo chunks phân đoạn
    if not chunks:
        chunks = []
        for p in pages_data:
            p_els = p["elements"]
            if not p_els:
                continue
            title_el = next((e for e in p_els if e["type"] in ("title", "h1", "h2")), p_els[0])
            cid = f"{doc_name}_slide_c{p['page_num']}"
            full_txt = "\n".join(e["text"] for e in p_els)
            chunks.append({
                "chunk_id": cid,
                "source": doc_name,
                "format": "pdf",
                "chunk_type": "pdf_slide",
                "breadcrumb": f"{doc_name} > Slide {p['page_num']}",
                "section_title": title_el["text"][:80],
                "content": full_txt,
                "raw_content": full_txt,
                "word_count": len(full_txt.split()),
                "char_length": len(full_txt),
                "page_start": p["page_num"],
                "page_end": p["page_num"],
                "page_numbers": [p["page_num"]],
                "element_ids": [e["element_id"] for e in p_els]
            })

    # Xác định tiêu đề văn bản thực tế bằng Visual DLA (0% Regex)
    doc_title = extract_document_title_dla(doc, pdf_path)
    if not doc_title or doc_title == doc_name or doc_title == pdf_path.stem:
        doc_title = next((e["text"] for e in all_elements if e["type"] == "title"), None) or clean_filename_fallback(pdf_path.stem)

    # 5. Ánh xạ 2 chiều giữa elements và chunks thông minh, triệt tiêu gối đầu và lộn xộn
    table_chunks_list = [c for c in chunks if c.get("chunk_type") in ("pdf_table", "pdf_vision_table", "docx_table")]
    text_chunks_list = [c for c in chunks if c.get("chunk_type") not in ("pdf_table", "pdf_vision_table", "docx_table")]

    for el in all_elements:
        if el["type"] == "header_footer":
            el["chunk_id"] = None
            continue

        el_txt = el["text"].strip()
        if not el_txt:
            continue

        assigned_cid = None
        # A. Nếu là phần tử bảng biểu -> Ưu tiên gán theo row bounding boxes và table_row_indices
        if el["type"] == "table":
            el_yc = (el["box"][1] + el["box"][3]) / 2.0
            matching_tcs = [tc for tc in table_chunks_list if el["page_num"] in (set(tc.get("page_numbers") or [tc.get("page_start", 0)]))]

            # 1. Thử so khớp theo cấu trúc hàng nếu chunk có table_row_indices
            p_idx = el["page_num"] - 1
            if 0 <= p_idx < len(doc):
                p_obj = doc[p_idx]
                tabs = p_obj.find_tables()
                if tabs and tabs.tables:
                    for t in tabs.tables:
                        for r_idx, r in enumerate(t.rows):
                            if r.bbox[1] - 3 <= el_yc <= r.bbox[3] + 3:
                                if r_idx == 0 and matching_tcs:
                                    assigned_cid = matching_tcs[0].get("chunk_id")
                                    break
                                for tc in matching_tcs:
                                    row_indices = tc.get("table_row_indices", [])
                                    if r_idx in row_indices:
                                        assigned_cid = tc.get("chunk_id")
                                        break
                                if assigned_cid:
                                    break
                        if assigned_cid:
                            break

            # 2. Nếu chưa gán được, fallback theo so khớp nội dung text
            if not assigned_cid:
                best_score = 0
                for tc in matching_tcs:
                    tc_content = tc.get("content", "")
                    tc_body = tc_content.split("Chi tiết từng hàng:\n", 1)[-1] if "Chi tiết từng hàng:\n" in tc_content else tc_content
                    if el_txt in tc_body:
                        score = len(el_txt)
                        if score > best_score:
                            best_score = score
                            assigned_cid = tc.get("chunk_id")

            if not assigned_cid and matching_tcs:
                assigned_cid = matching_tcs[0].get("chunk_id")
        else:
            # B. Nếu là phần tử text (title, body, footnote, v.v.)
            # Tuyệt đối KHÔNG gán phần tử Document Title / Header trên đỉnh trang 1 vào chunk
            is_doc_title_element = False
            if el["page_num"] == 1 and el["box"][1] < 220:
                norm_el = "".join(ch for ch in el_txt.lower() if ch.isalnum())
                norm_doc_t = "".join(ch for ch in doc_title.lower() if ch.isalnum()) if doc_title else ""
                if el["type"] in ("title", "h1", "header_footer"):
                    is_doc_title_element = True
                elif norm_doc_t and len(norm_el) >= 8 and (norm_el in norm_doc_t or norm_doc_t in norm_el):
                    is_doc_title_element = True
                elif el_txt.lower().startswith(("mẫu số", "mau so", "quyết định", "ban hành")):
                    is_doc_title_element = True

            if not is_doc_title_element:
                # Kiểm tra xem có phải là footnote gắn vào table chunk không
                for tc in table_chunks_list:
                    tc_pages = set(tc.get("page_numbers") or [tc.get("page_start", 0)])
                    if el["page_num"] in tc_pages:
                        tc_content = tc.get("content", "")
                        tc_body = tc_content.split("Chi tiết từng hàng:\n", 1)[-1] if "Chi tiết từng hàng:\n" in tc_content else tc_content
                        if f"[Ghi chú / Footnote]: {el_txt}" in tc_content or (len(el_txt) >= 15 and el_txt in tc_body):
                            assigned_cid = tc.get("chunk_id")
                            break

            # Nếu không phải footnote của table, tìm trong text chunks
            if not assigned_cid and not is_doc_title_element:
                best_score = 0
                for txt_c in text_chunks_list:
                    c_pages = set(txt_c.get("page_numbers") or [txt_c.get("page_start", 0)])
                    if el["page_num"] in c_pages or not c_pages:
                        c_text = txt_c.get("content", "") or txt_c.get("raw_content", "")
                        c_body = c_text.strip()
                        if c_body.startswith("[Breadcrumbs:"):
                            c_body = c_body.split("]\n", 1)[-1].strip()
                        c_stitle = txt_c.get("section_title", "")
                        c_bc = txt_c.get("breadcrumb", "")

                        score = 0
                        if el_txt in c_body:
                            score = 1000 + len(el_txt)
                            if c_body.startswith(el_txt) or c_body.startswith(f"### {el_txt}"):
                                score += 500
                        elif len(el_txt) >= 15 and el_txt[:25] in c_body:
                            score = 500 + len(el_txt[:25])
                        elif el_txt in c_bc or (len(el_txt) >= 5 and c_stitle and el_txt in c_stitle):
                            score = 200 + len(el_txt)
                        elif txt_c.get("chunk_type") == "pdf_ocr_scan" and el["page_num"] in c_pages:
                            score = 50 + len(el_txt)

                        if score > best_score:
                            best_score = score
                            assigned_cid = txt_c.get("chunk_id")

        el["chunk_id"] = assigned_cid
        if assigned_cid:
            chunk_to_elements.setdefault(assigned_cid, []).append(el["element_id"])

    doc.close()

    # Chuẩn hóa breadcrumb cho mọi chunks: triệt tiêu 100% tên file thô sơ và rác OCR
    stem_prefix = pdf_path.stem
    for c in chunks:
        cid = c.get("chunk_id")
        eids = chunk_to_elements.get(cid, [])
        if not eids:
            c_pages = set(c.get("page_numbers") or [c.get("page_start", 1)])
            cand_els = [el for el in all_elements if el["page_num"] in c_pages and el["type"] != "header_footer"]
            if not cand_els:
                cand_els = [el for el in all_elements if el["type"] != "header_footer"]
            if cand_els:
                eids = [cand_els[0]["element_id"]]
                cand_els[0]["chunk_id"] = cid
        c["element_ids"] = eids

        bc = c.get("breadcrumb", "")
        tokens = [t.strip() for t in bc.split(" > ") if t.strip()]
        if not tokens:
            c["breadcrumb"] = f"{doc_title} > {c.get('section_title') or f'Trang {c.get('page_start', 1)}'}"
        else:
            tokens[0] = doc_title
            filtered_tokens = [tokens[0]]
            for t in tokens[1:]:
                if t == doc_title or t.lower().startswith("bảng biểu trang"):
                    continue
                if len(t) >= 12 and not is_valid_title_candidate(t):
                    continue
                filtered_tokens.append(t)
            if len(filtered_tokens) == 1 and c.get("section_title") and c.get("section_title") != doc_title:
                st = str(c.get("section_title")).strip()
                if not st.lower().startswith("bảng biểu trang"):
                    filtered_tokens.append(st)
            c["breadcrumb"] = " > ".join(filtered_tokens)

        if not c.get("section_title") or str(c["section_title"]).strip() in ("None", ""):
            tokens_after = [t.strip() for t in c["breadcrumb"].split(" > ") if t.strip()]
            c["section_title"] = tokens_after[-1] if len(tokens_after) > 1 else (doc_title or "Tổng quan")

    return {
        "filename": doc_name,
        "format": "pdf",
        "title": doc_title,
        "total_pages": total_pages,
        "processed_pages": max_process_pages,
        "elements_count": len(all_elements),
        "chunks_count": len(chunks),
        "pages": pages_data,
        "chunks": chunks
    }


def process_docx_document(
    docx_path: Path,
    chunks: list[dict]
) -> dict:
    """
    Trích xuất layout elements cho tệp DOCX bằng Character-Weighted Font Histogram & AST (0% Regex).
    """
    doc = docx.Document(docx_path)
    doc_name = docx_path.name

    elements = []
    chunk_to_elements = {}

    # 1. Thu thập phân phối cỡ chữ (Font Size Histogram) theo trọng số ký tự
    font_sizes = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        p_size = None
        for run in p.runs:
            if run.font.size:
                p_size = run.font.size.pt
                break
        if p_size:
            font_sizes.extend([p_size] * len(txt))

    # Xác định cỡ chữ chuẩn
    if font_sizes:
        size_counts = Counter(font_sizes)
        body_font_size = size_counts.most_common(1)[0][0]
    else:
        body_font_size = 11.0

    # 2. Duyệt từng đoạn văn và phân cấp theo OpenXML AST Style & Font Histogram
    elem_idx = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue

        style_name = (p.style.name or "").lower()
        
        # Cỡ chữ đoạn văn
        p_size = body_font_size
        for run in p.runs:
            if run.font.size:
                p_size = run.font.size.pt
                break

        # Phân loại cấp bậc (0% Regex)
        if "title" in style_name or p_size >= body_font_size * 1.5:
            elem_type = "title"
        elif "heading 1" in style_name or (p_size >= body_font_size * 1.25 and len(txt) < 120):
            elem_type = "heading_1"
        elif "heading 2" in style_name or (p_size >= body_font_size * 1.12 and len(txt) < 140):
            elem_type = "heading_2"
        elif "heading 3" in style_name or (p_size > body_font_size and len(txt) < 100):
            elem_type = "heading_3"
        else:
            elem_type = "body"

        eid = f"{doc_name}_e{elem_idx}"
        elem_entry = {
            "element_id": eid,
            "page_num": 1 + (elem_idx // 20),  # Phân trang logic tương đối (20 đoạn/trang)
            "box": [50.0, float(elem_idx * 30), 550.0, float(elem_idx * 30 + 24)],
            "text": txt,
            "type": elem_type,
            "confidence": 0.98,
            "chunk_id": None
        }
        elements.append(elem_entry)
        elem_idx += 1

    # 3. Duyệt các bảng biểu (Table)
    for tbl_idx, tbl in enumerate(doc.tables):
        tbl_txt_lines = []
        for row in tbl.rows:
            row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_txt:
                tbl_txt_lines.append(row_txt)

        if tbl_txt_lines:
            eid = f"{doc_name}_tbl_{tbl_idx}"
            full_tbl_text = "\n".join(tbl_txt_lines)
            elem_entry = {
                "element_id": eid,
                "page_num": 1 + (elem_idx // 20),
                "box": [40.0, float(elem_idx * 30), 560.0, float(elem_idx * 30 + 80)],
                "text": full_tbl_text[:300],
                "type": "table",
                "confidence": 0.99,
                "chunk_id": None
            }
            elements.append(elem_entry)
            elem_idx += 1

    # 4. Ánh xạ 2 chiều giữa elements và chunks
    for c in chunks:
        cid = c.get("chunk_id")
        c_text = c.get("content", "") or c.get("raw_content", "")
        for el in elements:
            if el["text"] and (el["text"] in c_text or c_text in el["text"] or el["text"][:40] in c_text):
                el["chunk_id"] = cid
                chunk_to_elements.setdefault(cid, []).append(el["element_id"])

    for c in chunks:
        cid = c.get("chunk_id")
        eids = chunk_to_elements.get(cid, [])
        if not eids and elements:
            eids = [elements[0]["element_id"]]
            if not elements[0].get("chunk_id"):
                elements[0]["chunk_id"] = cid
        c["element_ids"] = eids

    # Phân nhóm elements theo trang logic
    pages_dict = {}
    for el in elements:
        p_num = el["page_num"]
        pages_dict.setdefault(p_num, []).append(el)

    pages_data = [
        {
            "page_num": p_num,
            "width": 595.0,
            "height": 842.0,
            "elements": p_els
        }
        for p_num, p_els in sorted(pages_dict.items())
    ]

    first_title = next((e["text"] for e in elements if e["type"] == "title"), None)
    if not first_title or len(first_title) < 5:
        first_title = clean_filename_fallback(docx_path.stem)

    for c in chunks:
        bc = c.get("breadcrumb", "")
        tokens = [t.strip() for t in bc.split(" > ") if t.strip()]
        if not tokens:
            c["breadcrumb"] = f"{first_title} > {c.get('section_title') or f'Trang {c.get('page_start', 1)}'}"
        else:
            tokens[0] = first_title
            filtered_tokens = [tokens[0]]
            for t in tokens[1:]:
                if t == first_title:
                    continue
                filtered_tokens.append(t)
            c["breadcrumb"] = " > ".join(filtered_tokens)

    return {
        "filename": doc_name,
        "format": "docx",
        "title": first_title,
        "total_pages": len(pages_data),
        "processed_pages": len(pages_data),
        "elements_count": len(elements),
        "chunks_count": len(chunks),
        "pages": pages_data,
        "chunks": chunks
    }


def main():
    print("=" * 72)
    print("   BẮT ĐẦU TRÍCH XUẤT LAYOUT 5D DLA & CHUNKS TOÀN BỘ 108 TỆP (0% REGEX)")
    print("=" * 72)

    chunks_grouped = load_all_chunks()
    print(f"[*] Đã nạp chunks của {len(chunks_grouped)} tài liệu từ all_chunks.json")

    pdf_files = sorted(list(PDF_DIR.glob("*.pdf")))
    docx_files = sorted(list(DOCX_DIR.glob("*.docx")))

    total_files = len(pdf_files) + len(docx_files)
    print(f"[*] Tổng số tệp phát hiện: {total_files} ({len(pdf_files)} PDF, {len(docx_files)} DOCX)")

    manifest = []
    processed_count = 0

    # 1. Xử lý 57 tệp PDF
    print("\n--- 1. XỬ LÝ 57 TỆP PDF NGÂN HÀNG & SLIDE ---")
    all_standardized_chunks = []
    for idx, pdf_path in enumerate(pdf_files, 1):
        fn = pdf_path.name
        # Kiểm tra xem tài liệu có lớp text số (digital PDF) hay scan ảnh thuần
        is_digital = False
        try:
            doc_peek = pymupdf.open(pdf_path)
            sample_txt_len = sum(len(p.get_text("text").strip()) for p in doc_peek[:min(3, len(doc_peek))])
            doc_peek.close()
            is_digital = (sample_txt_len > 80)
        except Exception:
            is_digital = False

        if is_digital:
            # Đối với digital PDF: sinh chunks động bằng PDFChunker mới nhất để loại bỏ 100% duplicate table và gắn footnote
            try:
                chunks = pdf_chunker_instance.chunk_pdf(pdf_path, max_pages=15)
            except Exception as ce:
                print(f"  [Chunker fallback {fn}]: {ce}")
                chunks = chunks_grouped.get(fn, [])
        else:
            chunks = chunks_grouped.get(fn, [])
        try:
            doc_data = process_pdf_document(pdf_path, chunks)
            out_file = OUTPUT_LAYOUTS_DIR / f"{fn}.json"
            with open(out_file, "w", encoding="utf-8") as f:
                json.dump(doc_data, f, ensure_ascii=False)

            all_standardized_chunks.extend(doc_data["chunks"])

            manifest.append({
                "id": fn,
                "name": fn,
                "format": "pdf",
                "title": doc_data["title"],
                "total_pages": doc_data["total_pages"],
                "elements_count": doc_data["elements_count"],
                "chunks_count": doc_data["chunks_count"],
                "raw_url": f"/dataset_raw/pdf/{fn}"
            })
            processed_count += 1
            if idx % 5 == 0 or idx == len(pdf_files) or "vcb_pdf_26" in fn:
                print(f"  [{idx}/{len(pdf_files)}] PDF: {fn} -> {doc_data['elements_count']} boxes, {doc_data['chunks_count']} chunks", flush=True)
        except Exception as e:
            print(f"  [!] Lỗi xử lý {fn}: {e}", flush=True)

    # 2. Xử lý 51 tệp DOCX
    print("\n--- 2. XỬ LÝ 51 TỆP DOCX QUY CHẾ & HƯỚNG DẪN ---", flush=True)
    for idx, docx_path in enumerate(docx_files, 1):
        fn = docx_path.name
        chunks = chunks_grouped.get(fn, [])
        try:
            doc_data = process_docx_document(docx_path, chunks)
            out_file = OUTPUT_LAYOUTS_DIR / f"{fn}.json"
            with open(out_file, "w", encoding="utf-8") as f:
                json.dump(doc_data, f, ensure_ascii=False)

            all_standardized_chunks.extend(doc_data["chunks"])

            manifest.append({
                "id": fn,
                "name": fn,
                "format": "docx",
                "title": doc_data["title"],
                "total_pages": doc_data["total_pages"],
                "elements_count": doc_data["elements_count"],
                "chunks_count": doc_data["chunks_count"],
                "raw_url": f"/dataset_raw/docx/{fn}"
            })
            processed_count += 1
            if idx % 10 == 0 or idx == len(docx_files):
                print(f"  [{idx}/{len(docx_files)}] DOCX: {fn} -> {doc_data['elements_count']} boxes, {doc_data['chunks_count']} chunks", flush=True)
        except Exception as e:
            print(f"  [!] Lỗi xử lý {fn}: {e}", flush=True)

    # 3. Đồng bộ toàn bộ chunks chuẩn hóa ra all_chunks.json
    for out_chunk_path in [CHUNK_CANDIDATE_PATHS[0], CHUNK_CANDIDATE_PATHS[2]]:
        out_chunk_path.parent.mkdir(parents=True, exist_ok=True)
        with open(out_chunk_path, "w", encoding="utf-8") as f:
            json.dump(all_standardized_chunks, f, ensure_ascii=False, indent=2)
    print(f"[*] Đã đồng bộ {len(all_standardized_chunks)} chunks chuẩn hóa (0% raw filename) vào {CHUNK_CANDIDATE_PATHS[0].name}")

    # 4. Ghi file Manifest
    with open(OUTPUT_MANIFEST_JSON, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

    with open(OUTPUT_MANIFEST_JS, "w", encoding="utf-8") as f:
        f.write("window.DOCUMENTS_108_MANIFEST = " + json.dumps(manifest, ensure_ascii=False, indent=2) + ";\n")

    print("\n" + "=" * 72)
    print(f"[OK] ĐÃ HOÀN TẤT XỬ LÝ {processed_count}/{total_files} TỆP BẰNG 0% REGEX DLA!")
    print(f"[*] Manifest JSON: {OUTPUT_MANIFEST_JSON}")
    print(f"[*] Manifest JS:   {OUTPUT_MANIFEST_JS}")
    print(f"[*] Layouts Dir:   {OUTPUT_LAYOUTS_DIR}")
    print("=" * 72)


if __name__ == "__main__":
    main()
